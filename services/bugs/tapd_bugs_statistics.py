"""
Bug统计报表生成工具（包含致命/严重/一般三级分类，支持致命+严重合并两级模式）
与用户提供的图片结构完全兼容
"""
import time

import pandas as pd
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from tapdsdk.sdk import TapdAPIClient
from datetime import datetime, date
from services.bugs.tapd_bugs_fetcher import TapdBugsFetcher
import xlwings as xw
import os
import yaml
import matplotlib.pyplot as plt

from services.util import utils


class TapdBugStatistics:
    def __init__(self, client_id: str, client_secret: str):
        self.client_id = client_id
        self.client_secret = client_secret

    def generate_statistics(self, version_report: str,
                            workspace_id: int,
                            severity_mode: str = "2level",
                            ) -> str:
        """生成分析报告"""

        # 将 client_id 和 client_secret 传递给 TapdBugsFetcher
        bugs_fetcher = TapdBugsFetcher(self.client_id, self.client_secret)
        bugs = bugs_fetcher.fetch_bugs(version_report, workspace_id)

        now = datetime.now()
        formatted_time = now.strftime("%Y-%m-%d %H:%M:%S")

        file = generate_statistics(
            api_data=bugs,
            title=f"{version_report}_{formatted_time}_bug统计报表",
            output_file=f"{version_report}_bug统计报表.xlsx",
            severity_mode=severity_mode
        )
        return file

def generate_statistics(api_data, title="bug统计报表", output_file="bug_report.xlsx", severity_mode="2level"):
    # ====== 输出目录 ======
    output_dir = "output_statistics"
    os.makedirs(output_dir, exist_ok=True)
    # 拼接输出文件路径
    output_file = os.path.join(output_dir, output_file)

    # ================== 数据预处理 ==================
    def preprocess_data(data):
        stats = defaultdict(
            lambda: {
                'non_auto': {
                    'unresolved': defaultdict(int),
                    'resolved': defaultdict(int)
                },
                'auto': {
                    'unresolved': defaultdict(int),
                    'resolved': defaultdict(int)
                }
            }
        )

        for ticket in data:
            module = str(ticket.get('module', '')).strip() or 'Other'
            title = ticket.get('title', '')
            is_auto = any(kw in title for kw in ['自动化', '自动提单', '接口测试'])

            status = 'resolved' if ticket.get('status') in ['已解决', '验证关闭', '排查关闭', '延期解决'] else 'unresolved'
            severity = ticket.get('severity', '一般').strip()
            if severity_mode == "3level":
                sev_key = 'fatal' if severity == '致命' else 'severe' if severity == '严重' else 'normal'
            else:
                sev_key = 'severe' if severity in ['致命', '严重'] else 'normal'

            key_type = 'auto' if is_auto else 'non_auto'
            stats[module][key_type][status][sev_key] += 1

        return stats

    # ================== 生成DataFrame ==================
    def create_dataframe(stats):
        rows = []
        for module in stats:
            non_auto = stats[module]['non_auto']
            auto = stats[module]['auto']

            def get_value(data, key):
                return data.get(key, 0) or 0  # 确保返回整数

            def format_num(num):
                return num if num != 0 else ''

            non_auto_total = sum(non_auto['unresolved'].values()) + sum(non_auto['resolved'].values())
            auto_total = sum(auto['unresolved'].values()) + sum(auto['resolved'].values())

            if severity_mode == "3level":
                unresolved = {
                    'fatal': get_value(non_auto['unresolved'], 'fatal') + get_value(auto['unresolved'], 'fatal'),
                    'severe': get_value(non_auto['unresolved'], 'severe') + get_value(auto['unresolved'], 'severe'),
                    'normal': get_value(non_auto['unresolved'], 'normal') + get_value(auto['unresolved'], 'normal')
                }
                resolved = {
                    'fatal': get_value(non_auto['resolved'], 'fatal') + get_value(auto['resolved'], 'fatal'),
                    'severe': get_value(non_auto['resolved'], 'severe') + get_value(auto['resolved'], 'severe'),
                    'normal': get_value(non_auto['resolved'], 'normal') + get_value(auto['resolved'], 'normal')
                }
                row = {
                    '模块': module,
                    '非自动化bug': format_num(non_auto_total),
                    '自动化bug': format_num(auto_total),
                    'Bug总数': format_num(non_auto_total + auto_total),
                    '未解决_致命': format_num(unresolved['fatal']),
                    '未解决_严重': format_num(unresolved['severe']),
                    '未解决_一般': format_num(unresolved['normal']),
                    '未解决_总计': format_num(sum(unresolved.values())),
                    '已解决_致命': format_num(resolved['fatal']),
                    '已解决_严重': format_num(resolved['severe']),
                    '已解决_一般': format_num(resolved['normal']),
                    '已解决_总计': format_num(sum(resolved.values())),
                    '致命总计': format_num(unresolved['fatal'] + resolved['fatal']),
                    '严重总计': format_num(unresolved['severe'] + resolved['severe']),
                    '一般总计': format_num(unresolved['normal'] + resolved['normal']),
                    '解决率_致命': _calc_rate(resolved['fatal'], unresolved['fatal'] + resolved['fatal']),
                    '解决率_严重': _calc_rate(resolved['severe'], unresolved['severe'] + resolved['severe']),
                    '解决率_一般': _calc_rate(resolved['normal'], unresolved['normal'] + resolved['normal'])
                }
            else:
                # 2列模式：致命+严重合并为严重
                unresolved = {
                    'severe': get_value(non_auto['unresolved'], 'severe') + get_value(auto['unresolved'], 'severe') +
                              get_value(non_auto['unresolved'], 'fatal') + get_value(auto['unresolved'], 'fatal'),
                    'normal': get_value(non_auto['unresolved'], 'normal') + get_value(auto['unresolved'], 'normal')
                }
                resolved = {
                    'severe': get_value(non_auto['resolved'], 'severe') + get_value(auto['resolved'], 'severe') +
                              get_value(non_auto['resolved'], 'fatal') + get_value(auto['resolved'], 'fatal'),
                    'normal': get_value(non_auto['resolved'], 'normal') + get_value(auto['resolved'], 'normal')
                }
                row = {
                    '模块': module,
                    '非自动化bug': format_num(non_auto_total),
                    '自动化bug': format_num(auto_total),
                    'Bug总数': format_num(non_auto_total + auto_total),
                    '未解决_严重': format_num(unresolved['severe']),
                    '未解决_一般': format_num(unresolved['normal']),
                    '未解决_总计': format_num(sum(unresolved.values())),
                    '已解决_严重': format_num(resolved['severe']),
                    '已解决_一般': format_num(resolved['normal']),
                    '已解决_总计': format_num(sum(resolved.values())),
                    '严重总计': format_num(unresolved['severe'] + resolved['severe']),
                    '一般总计': format_num(unresolved['normal'] + resolved['normal']),
                    '解决率_严重': _calc_rate(resolved['severe'], unresolved['severe'] + resolved['severe']),
                    '解决率_一般': _calc_rate(resolved['normal'], unresolved['normal'] + resolved['normal'])
                }
            rows.append(row)

        df = pd.DataFrame(rows)

        # 添加总计行（特殊处理空值）
        total_row = {}
        for col in df.columns:
            if col == '模块':
                total_row[col] = '总计'
            elif df[col].dtype == 'object':  # 文本列处理
                valid_values = [x for x in df[col] if isinstance(x, (int, float))]
                total_row[col] = sum(valid_values)
            else:
                total_row[col] = df[col].sum()

        # 重新计算解决率
        if severity_mode == "3level":
            total_row['解决率_致命'] = _calc_rate(total_row['已解决_致命'], total_row['致命总计'])
            total_row['解决率_严重'] = _calc_rate(total_row['已解决_严重'], total_row['严重总计'])
            total_row['解决率_一般'] = _calc_rate(total_row['已解决_一般'], total_row['一般总计'])
        else:
            total_row['解决率_严重'] = _calc_rate(total_row['已解决_严重'], total_row['严重总计'])
            total_row['解决率_一般'] = _calc_rate(total_row['已解决_一般'], total_row['一般总计'])

        df = pd.concat([df, pd.DataFrame([total_row])], ignore_index=True)
        return df

    def _calc_rate(numerator, denominator):
        try:
            numerator = int(numerator)
        except (TypeError, ValueError):
            numerator = 0
        try:
            denominator = int(denominator)
        except (TypeError, ValueError):
            denominator = 0
        if denominator == 0:
            return '100%'
        rate = (numerator / denominator) * 100
        return f"{round(rate)}%"

    # ================== Excel样式设置 ==================
    def apply_excel_style(writer, title=None):
        workbook = writer.book
        worksheet = workbook.active

        # ===== 基础样式 =====
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        center_alignment = Alignment(horizontal='center', vertical='center')

        # ===== 主标题 ===== (匹配图片的蓝色底纹)
        if severity_mode == "3level":
            worksheet.merge_cells('A1:R1')
        else:
            worksheet.merge_cells('A1:N1')
        title_cell = worksheet['A1']
        title_cell.value = title
        title_cell.font = Font(size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="5B9BD5", fill_type="solid")
        title_cell.alignment = center_alignment

        # ===== 列组颜色配置 ===== (精确匹配图片列组颜色)
        if severity_mode == "3level":
            column_groups = {
                'H': "DDEBF7",  # 未解决-浅蓝
                'L': "E2EFDA",  # 已解决-浅绿
                'P:R': "E2EFDA"  # 解决率-浅绿
            }
        else:
            column_groups = {
                'G': "DDEBF7",  # 未解决-浅蓝
                'J': "E2EFDA",  # 已解决-浅绿
                'M:N': "E2EFDA"  # 解决率-浅绿
            }

        # ===== 二级标题样式修改（移除边框） =====
        if severity_mode == "3level":
            group_headers = [
                ('E2:H2', '未解决', "DDEBF7"),
                ('I2:L2', '已解决', "E2EFDA"),
                ('M2:O2', 'bug总计', "FFF2CC"),
                ('P2:R2', '解决率', "E2EFDA")
            ]
        else:
            group_headers = [
                ('E2:G2', '未解决', "DDEBF7"),
                ('H2:J2', '已解决', "E2EFDA"),
                ('K2:L2', 'bug总计', "FFF2CC"),
                ('M2:N2', '解决率', "E2EFDA")
            ]

        for merge_range, title2, color in group_headers:
            cell = worksheet[merge_range.split(':')[0]]
            worksheet.merge_cells(merge_range)
            cell.value = title2
            cell.fill = PatternFill(start_color=color, fill_type="solid")
            cell.font = Font(bold=True, color="000000")
            cell.alignment = center_alignment

        # ===== 三级标题 ===== (完全复现图片文本)
        if severity_mode == "3level":
            headers = [
                ('A3', '模块', 'FFFFFF'), ('B3', '非自动化bug', 'FFFFFF'),
                ('C3', '自动化bug', 'FFFFFF'), ('D3', 'Bug总数', 'FFFFFF'),
                ('E3', '致命', 'DDEBF7'), ('F3', '严重', 'DDEBF7'), ('G3', '一般', 'DDEBF7'), ('H3', '总计', 'DDEBF7'),
                ('I3', '致命', 'E2EFDA'), ('J3', '严重', 'E2EFDA'), ('K3', '一般', 'E2EFDA'), ('L3', '总计', 'E2EFDA'),
                ('M3', '致命', 'FFF2CC'), ('N3', '严重', 'FFF2CC'), ('O3', '一般', 'FFF2CC'),
                ('P3', '致命bug', 'E2EFDA'), ('Q3', '严重bug', 'E2EFDA'), ('R3', '一般bug', 'E2EFDA')
            ]
        else:
            headers = [
                ('A3', '模块', 'FFFFFF'), ('B3', '非自动化bug', 'FFFFFF'),
                ('C3', '自动化bug', 'FFFFFF'), ('D3', 'Bug总数', 'FFFFFF'),
                ('E3', '严重', 'DDEBF7'), ('F3', '一般', 'DDEBF7'), ('G3', '总计', 'DDEBF7'),
                ('H3', '严重', 'E2EFDA'), ('I3', '一般', 'E2EFDA'), ('J3', '总计', 'E2EFDA'),
                ('K3', '严重', 'FFF2CC'), ('L3', '一般', 'FFF2CC'),
                ('M3', '严重bug', 'E2EFDA'), ('N3', '一般bug', 'E2EFDA')
            ]
        for cell_ref, text, color in headers:
            cell = worksheet[cell_ref]
            cell.value = text
            cell.fill = PatternFill(start_color=color, fill_type="solid")
            cell.font = Font(bold=True)
            cell.border = thin_border
            cell.alignment = center_alignment

        # 数据区样式
        if severity_mode == "3level":
            data_cols = 18
        else:
            data_cols = 14
        for row in worksheet.iter_rows(min_row=4, max_row=worksheet.max_row - 1, min_col=1, max_col=data_cols):
            for cell in row:
                cell.border = thin_border
                cell.alignment = center_alignment
                # 应用列组颜色
                applied = False
                for cols, color in column_groups.items():
                    if ':' in cols:
                        start, end = cols.split(':')
                    else:
                        start = end = cols
                    if start <= cell.column_letter <= end:
                        cell.fill = PatternFill(start_color=color, fill_type="solid")
                        applied = True
                        break
                if not applied:
                    cell.fill = PatternFill(start_color="FFFFFF", fill_type="solid")

        # 合并单元格
        if severity_mode == "3level":
            merge_config = [
                ('A2:A3', '模块'),
                ('B2:B3', '非自动化bug'),
                ('C2:C3', '自动化bug'),
                ('D2:D3', 'Bug总数')
            ]
        else:
            merge_config = [
                ('A2:A3', '模块'),
                ('B2:B3', '非自动化bug'),
                ('C2:C3', '自动化bug'),
                ('D2:D3', 'Bug总数')
            ]
        for cell_range, text in merge_config:
            worksheet.merge_cells(cell_range)
            cell = worksheet[cell_range.split(':')[0]]
            cell.value = text
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.fill = PatternFill(start_color="FFFFFF", fill_type="solid")
        worksheet.freeze_panes = 'A4'

        # ===== 前三标记逻辑 =====
        def mark_top3(col_letter, start_row=4):
            color_map = [
                ('E60000', 1),  # 深红-第一
                ('FF6666', 2),  # 亮红-第二
                ('FFCC00', 3)   # 橙黄-第三
            ]
            values = []
            for row in range(start_row, worksheet.max_row):
                cell = worksheet[f"{col_letter}{row}"]
                try:
                    values.append(float(cell.value) if cell.value not in (None, "") else 0)
                except Exception:
                    values.append(0)
            sorted_values = sorted(set(values), reverse=True)[:3]
            for idx, val in enumerate(sorted_values):
                for row in range(start_row, worksheet.max_row):
                    cell = worksheet[f"{col_letter}{row}"]
                    try:
                        if float(cell.value) == val and val > 0:
                            cell.fill = PatternFill(start_color=color_map[idx][0], fill_type="solid")
                    except Exception:
                        continue

        if severity_mode == "3level":
            top3_cols = ['D', 'H', 'M', 'N', 'O']
            rate_cols = ['P', 'Q', 'R']
        else:
            top3_cols = ['D', 'G', 'K', 'L']
            rate_cols = ['M', 'N']

        for col in top3_cols:
            mark_top3(col)

        # ===== 解决率非100%标红 =====
        light_red = PatternFill(start_color="FFCCCC", fill_type="solid")
        for col in rate_cols:
            for row in range(4, worksheet.max_row):
                cell = worksheet[f"{col}{row}"]
                if cell.value:
                    clean_value = str(cell.value).replace('%', '')
                    try:
                        if float(clean_value) < 100:
                            cell.fill = light_red
                    except ValueError:
                        pass

        # 总计行样式
        total_row = worksheet.max_row
        yellow_fill = PatternFill(start_color="FFFF00", fill_type="solid")
        for cell in worksheet[total_row]:
            cell.fill = yellow_fill
            cell.font = Font(bold=True)
            cell.border = thin_border
            cell.alignment = center_alignment

        # 列宽
        if severity_mode == "3level":
            column_widths = {
                'A': 18, 'B': 12, 'C': 10, 'D': 9,
                'E': 6, 'F': 6, 'G': 6, 'H': 6,
                'I': 6, 'J': 6, 'K': 6, 'L': 6,
                'M': 6, 'N': 6, 'O': 6,
                'P': 9, 'Q': 9, 'R': 9
            }
        else:
            column_widths = {
                'A': 18, 'B': 12, 'C': 10, 'D': 9,
                'E': 6, 'F': 6, 'G': 6,
                'H': 6, 'I': 6, 'J': 6,
                'K': 6, 'L': 6,
                'M': 9, 'N': 9
            }
        for col, width in column_widths.items():
            worksheet.column_dimensions[col].width = width

        return writer

    # ================== 主流程 ==================
    try:
        stats = preprocess_data(api_data)
        df = create_dataframe(stats)

        # 列顺序（与图片完全一致）
        if severity_mode == "3level":
            columns_order = [
                '模块', '非自动化bug', '自动化bug', 'Bug总数',
                '未解决_致命', '未解决_严重', '未解决_一般', '未解决_总计',
                '已解决_致命', '已解决_严重', '已解决_一般', '已解决_总计',
                '致命总计', '严重总计', '一般总计',
                '解决率_致命', '解决率_严重', '解决率_一般'
            ]
        else:
            columns_order = [
                '模块', '非自动化bug', '自动化bug', 'Bug总数',
                '未解决_严重', '未解决_一般', '未解决_总计',
                '已解决_严重', '已解决_一般', '已解决_总计',
                '严重总计', '一般总计',
                '解决率_严重', '解决率_一般'
            ]
        df = df[columns_order]

        # 写入Excel
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, startrow=3, header=False)
            apply_excel_style(writer, title)

        return output_file
    except Exception as e:
        raise RuntimeError(f"报表生成失败: {str(e)}")

# 测试用例
if __name__ == "__main__":
    sample_data = [
        {'module': 'Audio', 'title': '噪音问题', 'status': '已解决', 'severity': '致命'},
        {'module': 'Audio', 'title': '杂音问题', 'status': '已解决', 'severity': '严重'},
        {'module': 'Audio', 'title': '[自动化]电流声', 'status': '未解决', 'severity': '一般'},
    ]
    try:
        output = generate_statistics(sample_data, severity_mode="3level")
        print(f"三列模式生成成功：{output}")
        output2 = generate_statistics(sample_data, severity_mode="2level")
        print(f"两列模式生成成功：{output2}")
    except Exception as e:
        print(f"生成失败：{str(e)}")