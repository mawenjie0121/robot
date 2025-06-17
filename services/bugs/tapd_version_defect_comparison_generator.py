from openpyxl.styles import PatternFill, Alignment, Border, Side, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from services.bugs.tapd_bugs_fetcher import TapdBugsFetcher


def blank_if_zero(val):
    return "" if val == 0 else val

class TapdVersionDefectComparisonGenerator:
    def __init__(self):
        self.version_data_dict = {}
        self.stats = {}
        self.all_modules = set()

    def analyze(self):
        for version, api_data in self.version_data_dict.items():
            stats = {}
            stats['总计'] = {
                'manual_total': 0, 'manual_fatal_serious': 0, 'manual_normal': 0, 'manual_crash': 0,
                'auto_total': 0, 'auto_fatal_serious': 0, 'auto_normal': 0, 'auto_crash': 0
            }
            for defect in api_data:
                original_module = defect.get('module', '')
                if not original_module:
                    original_module = '-空-'
                title = defect.get('title', '')
                severity = defect.get('severity', '')
                custom_value = defect.get('custom_field_33', '')

                if original_module == '插件':
                    if custom_value != '':
                        target_module = '插件-' + custom_value
                    else:
                        target_module = "插件--空-"
                else:
                    target_module = original_module

                self.all_modules.add(target_module)

                if target_module not in stats:
                    stats[target_module] = {
                        'manual_total': 0, 'manual_fatal_serious': 0, 'manual_normal': 0, 'manual_crash': 0,
                        'auto_total': 0, 'auto_fatal_serious': 0, 'auto_normal': 0, 'auto_crash': 0
                    }

                is_auto = any(kw in title for kw in ["自动化", "接口测试", "自动提单"])
                is_crash = "crash" in title.lower()

                if is_auto:
                    stats[target_module]['auto_total'] += 1
                    stats['总计']['auto_total'] += 1
                    if severity in ['致命', '严重']:
                        stats[target_module]['auto_fatal_serious'] += 1
                        stats['总计']['auto_fatal_serious'] += 1
                    if severity == '一般':
                        stats[target_module]['auto_normal'] += 1
                        stats['总计']['auto_normal'] += 1
                    if is_crash:
                        stats[target_module]['auto_crash'] += 1
                        stats['总计']['auto_crash'] += 1
                else:
                    stats[target_module]['manual_total'] += 1
                    stats['总计']['manual_total'] += 1
                    if severity in ['致命', '严重']:
                        stats[target_module]['manual_fatal_serious'] += 1
                        stats['总计']['manual_fatal_serious'] += 1
                    if severity == '一般':
                        stats[target_module]['manual_normal'] += 1
                        stats['总计']['manual_normal'] += 1
                    if is_crash:
                        stats[target_module]['manual_crash'] += 1
                        stats['总计']['manual_crash'] += 1
            self.stats[version] = stats

    def generate_excel(self, output_file='缺陷统计对比报表.xlsx'):
        if not self.stats:
            self.analyze()

        wb = Workbook()
        ws = wb.active
        ws.title = "缺陷对比"

        versions = list(self.version_data_dict.keys())
        module_list = sorted(self.all_modules)

        # 插件模块和其他模块分开
        plugin_modules = [m for m in module_list if m.startswith('插件-')]
        other_modules = [m for m in module_list if not m.startswith('插件-')]

        # 插件模块排序，把“插件--空-”放到最后
        plugin_modules_sorted = [m for m in plugin_modules if m != '插件--空-']
        if '插件--空-' in plugin_modules:
            plugin_modules_sorted.append('插件--空-')

        # 非插件模块排序，把“-空-”放到最后
        other_modules_sorted = [m for m in other_modules if m != '-空-']
        if '-空-' in other_modules:
            other_modules_sorted.append('-空-')

        # ====== 表头 ======
        ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=1)
        ws.merge_cells(start_row=1, start_column=2, end_row=3, end_column=2)
        ws['A1'] = '分组'
        ws['B1'] = '模块'
        ws['A1'].alignment = ws['B1'].alignment = Alignment(horizontal='center', vertical='center')
        ws['A1'].font = ws['B1'].font = Font(size=14, bold=True)

        col = 3
        for version in versions:
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+7)
            ws.cell(row=1, column=col, value=version)
            ws.cell(row=1, column=col).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=1, column=col).font = Font(size=14, bold=True)

            ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col+3)
            ws.merge_cells(start_row=2, start_column=col+4, end_row=2, end_column=col+7)
            ws.cell(row=2, column=col, value='手工')
            ws.cell(row=2, column=col+4, value='自动化')
            ws.cell(row=2, column=col).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=2, column=col+4).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=2, column=col).font = Font(size=13, bold=True)
            ws.cell(row=2, column=col+4).font = Font(size=13, bold=True)

            headers = ['总数', '致命+严重', '一般', 'crash数', '总数', '致命+严重', '一般', 'crash数']
            for i, h in enumerate(headers):
                cell = ws.cell(row=3, column=col+i, value=h)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.font = Font(size=12, bold=True)
                if h == '致命+严重':
                    cell.fill = PatternFill(start_color='FFC7CE', fill_type='solid')
                elif h == 'crash数':
                    cell.fill = PatternFill(start_color='C6EFCE', fill_type='solid')
                elif h == '一般':
                    cell.fill = PatternFill(start_color='FFFFCC99', fill_type='solid')
            col += 8

        # 冻结前两列和前三行
        ws.freeze_panes = 'C4'

        # ====== 数据行 ======
        row_idx = 4
        plugin_start = None
        plugin_end = None

        # 插件区
        plugin_row_map = {}  # 新增：插件模块名->行号
        if plugin_modules_sorted:
            plugin_start = row_idx
            for module in plugin_modules_sorted:
                plugin_name = module.replace('插件-', '', 1)
                if not plugin_name or plugin_name == '-空-':
                    plugin_name = '-空-'
                ws.cell(row=row_idx, column=1, value='插件')
                ws.cell(row=row_idx, column=2, value=plugin_name)
                ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(row=row_idx, column=1).font = Font(size=12)
                ws.cell(row=row_idx, column=2).font = Font(size=12)
                col = 3
                for version in versions:
                    data = self.stats[version].get(module, {
                        'manual_total': 0, 'manual_fatal_serious': 0, 'manual_normal': 0, 'manual_crash': 0,
                        'auto_total': 0, 'auto_fatal_serious': 0, 'auto_normal': 0, 'auto_crash': 0
                    })
                    ws.cell(row=row_idx, column=col,   value=blank_if_zero(data['manual_total']))
                    ws.cell(row=row_idx, column=col+1, value=blank_if_zero(data['manual_fatal_serious']))
                    ws.cell(row=row_idx, column=col+2, value=blank_if_zero(data['manual_normal']))
                    ws.cell(row=row_idx, column=col+3, value=blank_if_zero(data['manual_crash']))
                    ws.cell(row=row_idx, column=col+4, value=blank_if_zero(data['auto_total']))
                    ws.cell(row=row_idx, column=col+5, value=blank_if_zero(data['auto_fatal_serious']))
                    ws.cell(row=row_idx, column=col+6, value=blank_if_zero(data['auto_normal']))
                    ws.cell(row=row_idx, column=col+7, value=blank_if_zero(data['auto_crash']))
                    for i in range(8):
                        ws.cell(row=row_idx, column=col+i).alignment = Alignment(horizontal='center', vertical='center')
                        ws.cell(row=row_idx, column=col+i).font = Font(size=12)
                    col += 8
                plugin_row_map[module] = row_idx  # 记录插件模块行号
                row_idx += 1
            plugin_end = row_idx - 1
            ws.merge_cells(start_row=plugin_start, start_column=1, end_row=plugin_end, end_column=1)

        # 非插件区
        other_row_map = {}  # 新增：非插件模块名->行号
        for module in other_modules_sorted:
            display_name = module if module else '-空-'
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
            ws.cell(row=row_idx, column=1, value=display_name)
            ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=row_idx, column=1).font = Font(size=12)
            col = 3
            for version in versions:
                data = self.stats[version].get(module, {
                    'manual_total': 0, 'manual_fatal_serious': 0, 'manual_normal': 0, 'manual_crash': 0,
                    'auto_total': 0, 'auto_fatal_serious': 0, 'auto_normal': 0, 'auto_crash': 0
                })
                ws.cell(row=row_idx, column=col,   value=blank_if_zero(data['manual_total']))
                ws.cell(row=row_idx, column=col+1, value=blank_if_zero(data['manual_fatal_serious']))
                ws.cell(row=row_idx, column=col+2, value=blank_if_zero(data['manual_normal']))
                ws.cell(row=row_idx, column=col+3, value=blank_if_zero(data['manual_crash']))
                ws.cell(row=row_idx, column=col+4, value=blank_if_zero(data['auto_total']))
                ws.cell(row=row_idx, column=col+5, value=blank_if_zero(data['auto_fatal_serious']))
                ws.cell(row=row_idx, column=col+6, value=blank_if_zero(data['auto_normal']))
                ws.cell(row=row_idx, column=col+7, value=blank_if_zero(data['auto_crash']))
                for i in range(8):
                    ws.cell(row=row_idx, column=col+i).alignment = Alignment(horizontal='center', vertical='center')
                    ws.cell(row=row_idx, column=col+i).font = Font(size=12)
                col += 8
            other_row_map[module] = row_idx  # 记录非插件模块行号
            row_idx += 1

        # 记录总计行号
        total_row = row_idx

        # 总计行
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
        ws.cell(row=row_idx, column=1, value='总计')
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=row_idx, column=1).font = Font(size=12, bold=True)
        col = 3
        for version in versions:
            data = self.stats[version].get('总计', {
                'manual_total': 0, 'manual_fatal_serious': 0, 'manual_normal': 0, 'manual_crash': 0,
                'auto_total': 0, 'auto_fatal_serious': 0, 'auto_normal': 0, 'auto_crash': 0
            })
            ws.cell(row=row_idx, column=col,   value=blank_if_zero(data['manual_total']))
            ws.cell(row=row_idx, column=col+1, value=blank_if_zero(data['manual_fatal_serious']))
            ws.cell(row=row_idx, column=col+2, value=blank_if_zero(data['manual_normal']))
            ws.cell(row=row_idx, column=col+3, value=blank_if_zero(data['manual_crash']))
            ws.cell(row=row_idx, column=col+4, value=blank_if_zero(data['auto_total']))
            ws.cell(row=row_idx, column=col+5, value=blank_if_zero(data['auto_fatal_serious']))
            ws.cell(row=row_idx, column=col+6, value=blank_if_zero(data['auto_normal']))
            ws.cell(row=row_idx, column=col+7, value=blank_if_zero(data['auto_crash']))
            for i in range(8):
                ws.cell(row=row_idx, column=col+i).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(row=row_idx, column=col+i).font = Font(size=12, bold=True)
            col += 8

        # ====== 新增：插件区后每5行填充颜色 ======
        color_fills = [
            PatternFill(start_color='FFF2F2F2', end_color='FFF2F2F2', fill_type='solid'),  # 浅灰
            PatternFill(start_color='FFEAF6FB', end_color='FFEAF6FB', fill_type='solid'),  # 浅蓝
            PatternFill(start_color='FFFFFBE5', end_color='FFFFFBE5', fill_type='solid'),  # 浅黄
            PatternFill(start_color='FFF6F0FA', end_color='FFF6F0FA', fill_type='solid'),  # 浅紫
        ]
        if plugin_end is not None:
            fill_start_row = plugin_end + 1
            fill_end_row = total_row - 1  # 不包含总计行
            if fill_start_row <= fill_end_row:
                for i, row in enumerate(range(fill_start_row, fill_end_row + 1)):
                    fill_idx = (i // 5) % 4
                    fill = color_fills[fill_idx]
                    for col in range(1, 2 + 8 * len(versions) + 1):
                        ws.cell(row=row, column=col).fill = fill

        # ====== 列宽和行高 ======
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 28
        for i in range(3, 3 + 8 * len(versions)):
            ws.column_dimensions[get_column_letter(i)].width = 10
        for row in range(1, ws.max_row + 1):
            ws.row_dimensions[row].height = 20

        thin = Side(border_style="thin", color="000000")
        medium = Side(border_style="medium", color="000000")

        max_col = 2 + 8 * len(versions)
        max_row = ws.max_row

        # 1. 表头区域（第1~3行）加细线边框
        for row in ws.iter_rows(min_row=1, max_row=3, min_col=1, max_col=max_col):
            for cell in row:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        # 2. 数据区域
        for row in range(4, max_row + 1):
            for col in range(1, max_col + 1):
                left = None
                right = None
                top = None
                bottom = None

                # 模块分组列（第1、2列）加细线四边
                if col in [1, 2]:
                    left = thin
                    right = thin
                    top = thin
                # 每个版本分隔列的左侧加medium竖线
                if col in [2 + i * 8 + 1 for i in range(1, len(versions))]:
                    left = medium
                # 最右侧一列加右边框
                if col == max_col:
                    right = thin
                # 最后一行加下边框
                if row == max_row:
                    bottom = thin

                # 只在有边框的单元格设置
                if left or right or top or bottom:
                    ws.cell(row=row, column=col).border = Border(left=left, right=right, top=top, bottom=bottom)

        # 定义颜色
        red_fill = PatternFill(start_color='DA0101', end_color='DA0101', fill_type='solid')
        white_font = Font(color='FFFFFF')
        green_fill = PatternFill(start_color='00B050', end_color='00B050', fill_type='solid')  # 绿色

        # 计算所有模块（插件+非插件）和总计行的行号
        module_row_map = {}
        module_row_map.update(plugin_row_map)
        module_row_map.update(other_row_map)
        # 总计行
        total_row = ws.max_row
        module_row_map['总计'] = total_row

        # 需要对比的手工列偏移（总数、致命+严重、一般、crash数）
        manual_compare_offsets = [0, 1, 2, 3]

        # 对每个模块（包括插件模块）和总计行，做手工相关单元格的对比
        for module, m_row in module_row_map.items():
            for v_idx in range(1, len(versions)):
                base_col = 3 + v_idx * 8  # 每个版本的手工总数列
                prev_base_col = base_col - 8
                for offset in manual_compare_offsets:
                    col = base_col + offset
                    prev_col = prev_base_col + offset
                    curr_val = ws.cell(row=m_row, column=col).value
                    prev_val = ws.cell(row=m_row, column=prev_col).value
                    # 空字符串转0
                    curr_val = int(curr_val) if curr_val not in (None, "") else 0
                    prev_val = int(prev_val) if prev_val not in (None, "") else 0
                    if curr_val == 0:
                        continue  # 当前为0不处理
                    if curr_val > prev_val:
                        ws.cell(row=m_row, column=col).fill = red_fill
                        ws.cell(row=m_row, column=col).font = white_font
                    elif curr_val < prev_val:
                        ws.cell(row=m_row, column=col).fill = green_fill

        wb.save(output_file)




        return

    def generate_defect_report(self, workspace_id, version_list, output_file='缺陷统计对比报表.xlsx'):
        bugs_fetcher = TapdBugsFetcher()
        self.version_data_dict = {}
        for version in version_list:
            api_data = bugs_fetcher.fetch_bugs(version, workspace_id, severity="fatal|normal|serious|prompt|advice")
            self.version_data_dict[version] = api_data
        self.stats = {}
        self.all_modules = set()
        return self.run(output_file=output_file)

    def run(self, output_file='缺陷统计对比报表.xlsx'):
        self.analyze()
        return self.generate_excel(output_file=output_file)

# 用法示例
if __name__ == '__main__':
    generator = TapdVersionDefectComparisonGenerator()
    workspace_id = 20428244
    version_list = [
        "liteAV 12.4（2.0）",
        "liteAV 12.5（2.0）",
        "liteAV 12.6（2.0）",
        # ... 你可以继续加
    ]
    generator.generate_defect_report(workspace_id, version_list, output_file='缺陷统计对比报表.xlsx')